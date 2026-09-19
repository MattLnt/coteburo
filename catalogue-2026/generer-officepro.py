#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Génère l'onglet OFFICEPRO de l'Excel catalogue Côté BURO.

Une ligne par déclinaison : une référence du tarif, c'est-à-dire un produit
dans un coloris.

Sources, dans le même dossier que ce script :
    officepro-references.json        le tarif extrait (parser-officepro.py)
    noms-officepro.json              le regroupement en fiches, le périmètre
    catalogue-officepro-pages.json   ce que le catalogue dit de chaque
                                     référence (parser-catalogue-officepro.py)

Le nom vient du catalogue, jamais du tarif. Le tarif nomme mal : il appelle
« CHAISE VERANO LOUNGE » ce que le catalogue titre « FAUTEUIL LOUNGE », et
« TECSEAT ETUDIANT » ce qu'il nomme « TECSEAT LEARNING ». Ces noms sont
arrêtés dans noms-officepro.json, chacun vérifié contre sa page.

Les colonnes de cotes restent vides : le catalogue OfficePro ne porte pas
de tableau de dimensions, seulement des légendes posées près des dessins,
que rien ne rattache à une référence. Voir _cotes dans noms-officepro.json.

Sortie :
    catalogue-officepro.xlsx

Usage :
    pip install openpyxl
    python generer-officepro.py
"""

import json
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

RACINE = Path(__file__).parent
REFERENCES = RACINE / "officepro-references.json"
NOMS = RACINE / "noms-officepro.json"
CATALOGUE = RACINE / "catalogue-officepro-pages.json"
SORTIE = RACINE / "catalogue-officepro.xlsx"

COLONNES = [
    ("Gamme", 18), ("Nature", 11), ("Sur devis", 10),
    ("Catégorie", 22), ("Sous-catégorie 1", 28), ("Sous-catégorie 2", 24),
    ("Nom sur le site", 46), ("Désignation fournisseur", 50),
    ("Page cat.", 10),
    ("Réf. produit", 13), ("Réf. finition", 30), ("Réf. complète", 18),
    ("Tarif HT", 10), ("Écotaxe", 9),
    ("Description générale", 46), ("Description technique", 52),
    ("Options (réf.)", 26),
    ("Larg. min", 10), ("Larg. max", 10),
    ("Prof. min", 10), ("Prof. max", 10),
    ("Haut. min", 10), ("Haut. max", 10),
    ("Poids (kg)", 10), ("Code EAN", 15),
]

MULTILIGNE = (15, 16)


def entetes(ws):
    fond = PatternFill("solid", fgColor="23262A")
    police = Font(color="FFFFFF", bold=True, size=10)
    for i, (titre, largeur) in enumerate(COLONNES, 1):
        c = ws.cell(row=1, column=i, value=titre)
        c.fill = fond
        c.font = police
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = largeur
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLONNES))}1"


def main():
    noms = json.loads(NOMS.read_text(encoding="utf-8"))
    gammes = {k: v for k, v in noms.items() if not k.startswith("_")}
    refs = {r["reference"]: r for r in json.loads(REFERENCES.read_text(encoding="utf-8"))}

    out = Workbook()
    ws = out.active
    ws.title = "OFFICEPRO"
    entetes(ws)

    ligne = 2
    stats = Counter()
    manquantes, sans_prix = [], []

    for gamme, conf in gammes.items():
        affichee = conf.get("gamme_affichee", gamme)
        cat = conf.get("categorie", "")
        sous = conf.get("sous_categories", [])
        devis = "Oui" if conf.get("sur_devis") else "Non"
        # Le texte de gamme ne sert plus que de repli : chaque fiche a le
        # sien. Servir la même description sur les vingt fiches d'une gamme
        # revenait à publier vingt pages en double.
        desc_gamme = conf.get("description", "")
        tech_gamme = conf.get("technique", [])
        opts = ", ".join(conf.get("options_refs", []))

        for produit in conf["produits"]:
            stats["fiches"] += 1
            nom = f"{produit['nom']} - {affichee}"
            # Le classement suit la gamme, sauf exception écrite sur la fiche.
            # Les coussins Arco sont fournis avec la banquette et ne se
            # vendent pas seuls : ils n'ont rien à faire au rayon Sièges.
            cat_fiche = produit.get("categorie", cat)
            sous_fiche = produit.get("sous_categories", sous)
            desc = produit.get("description") or desc_gamme
            tech = "\n".join(produit.get("technique") or tech_gamme)
            premiere = True

            for ref in produit["refs"]:
                r = refs.get(ref)
                if not r:
                    manquantes.append((gamme, ref))
                    continue
                if r["prix"] is None:
                    sans_prix.append(ref)

                vals = [
                    gamme, produit["nature"], devis, cat_fiche,
                    sous_fiche[0] if sous_fiche else "",
                    sous_fiche[1] if len(sous_fiche) > 1 else "",
                    nom, r["designation"], r["page_catalogue"] or "",
                    r["racine"], r["coloris"], ref,
                    r["prix"] if r["prix"] is not None else "",
                    r["ecotaxe"] if r["ecotaxe"] is not None else "",
                    desc if premiere else "",
                    tech if premiere else "",
                    opts if premiere else "",
                    # Cotes : le catalogue OfficePro n'en donne pas
                    # d'exploitable, on ne les invente pas.
                    "", "", "", "", "", "",
                    "", r["ean"],
                ]
                for j, v in enumerate(vals, 1):
                    ws.cell(row=ligne, column=j, value=v)
                for j in MULTILIGNE:
                    ws.cell(row=ligne, column=j).alignment = Alignment(
                        wrap_text=True, vertical="top"
                    )
                ligne += 1
                stats["lignes"] += 1
                premiere = False

    out.save(SORTIE)

    print(f"{stats['lignes']} lignes · {stats['fiches']} fiches · {len(gammes)} gammes")
    print(f"  sans prix (coussins inclus au produit) : {len(sans_prix)}")
    if manquantes:
        print(f"\n  {len(manquantes)} références introuvables au tarif :")
        for g, r in manquantes[:10]:
            print(f"     {g} · {r}")
    print(f"\nÉcrit → {SORTIE.name}")


if __name__ == "__main__":
    main()
