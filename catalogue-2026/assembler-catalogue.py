#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Assemble les trois onglets fournisseurs dans un seul classeur.

Chaque fournisseur a son générateur et son vérificateur, parce que chaque
tarif se lit autrement : Buronomic donne une ligne par finition, OfficePro
une ligne par coloris, Sokoa une ligne par catégorie de revêtement. Ce
script ne refait aucun de ces choix — il recopie les trois classeurs déjà
produits et contrôlés, dans l'ordre où le client les a listés.

Il vérifie au passage que les trois onglets parlent bien le même langage :
mêmes colonnes, dans le même ordre. Un onglet qui aurait dérivé serait
refusé ici plutôt que de partir à l'import.

Sources, dans le même dossier :
    catalogue-buronomic.xlsx
    catalogue-officepro.xlsx
    catalogue-sokoa.xlsx

Sortie :
    catalogue-coteburo.xlsx

Usage :
    python assembler-catalogue.py
"""

import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

RACINE = Path(__file__).parent
SORTIE = RACINE / "catalogue-coteburo.xlsx"

ONGLETS = [
    ("BURONOMIC", "catalogue-buronomic.xlsx"),
    ("OFFICEPRO", "catalogue-officepro.xlsx"),
    ("SOKOA", "catalogue-sokoa.xlsx"),
]

# Colonnes dont le contenu tient sur plusieurs lignes.
MULTILIGNE = (15, 16)


def main():
    out = Workbook()
    out.remove(out.active)

    reference = None
    total = 0
    resume = []

    for nom, fichier in ONGLETS:
        chemin = RACINE / fichier
        if not chemin.exists():
            print(f"✗ {fichier} manquant — lancer son générateur d'abord")
            return 1

        source = load_workbook(chemin, read_only=True).active
        lignes = list(source.iter_rows(values_only=True))
        entete = list(lignes[0])

        if reference is None:
            reference = entete
        elif entete != reference:
            print(f"✗ {nom} n'a pas les mêmes colonnes que {ONGLETS[0][0]} :")
            for i, (a, b) in enumerate(zip(reference, entete), 1):
                if a != b:
                    print(f"     colonne {i} : « {a} » vs « {b} »")
            return 1

        ws = out.create_sheet(nom)
        fond = PatternFill("solid", fgColor="23262A")
        police = Font(color="FFFFFF", bold=True, size=10)
        for j, titre in enumerate(entete, 1):
            c = ws.cell(row=1, column=j, value=titre)
            c.fill = fond
            c.font = police
            c.alignment = Alignment(vertical="center", wrap_text=True)

        # Les largeurs de colonne viennent de l'onglet d'origine : le nom
        # d'une fiche Sokoa est plus long que celui d'une fiche OfficePro.
        for lettre, dim in load_workbook(chemin).active.column_dimensions.items():
            if dim.width:
                ws.column_dimensions[lettre].width = dim.width

        for i, ligne in enumerate(lignes[1:], start=2):
            for j, valeur in enumerate(ligne, 1):
                ws.cell(row=i, column=j, value=valeur)
            for j in MULTILIGNE:
                ws.cell(row=i, column=j).alignment = Alignment(
                    wrap_text=True, vertical="top"
                )

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(entete))}1"

        n = len(lignes) - 1
        gammes = len({r[0] for r in lignes[1:] if r[0]})
        fiches = len({(r[0], r[6]) for r in lignes[1:] if r[6]})
        total += n
        resume.append((nom, n, fiches, gammes))

    out.save(SORTIE)

    largeur = max(len(n) for n, *_ in resume)
    print(f"{len(resume)} onglets · {len(reference)} colonnes\n")
    for nom, n, fiches, gammes in resume:
        print(f"  {nom:{largeur}}  {n:6} lignes  {fiches:4} fiches  {gammes:3} gammes")
    print(f"  {'':{largeur}}  {total:6} lignes au total")
    print(f"\nÉcrit → {SORTIE.name}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
