#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Génère l'onglet BURONOMIC de l'Excel catalogue Côté BURO.

Une ligne par déclinaison : un code article du tarif, c'est-à-dire un
produit dans une finition, déduction faite de ce que la rationalisation
2026 retire.

Sources, dans le même dossier que ce script :
    buronomic.xlsx              le tarif Buronomic 2026
    noms-produits.json          le regroupement des références en fiches,
                                le périmètre et les règles de rationalisation
    catalogue-pages-cotes.json  page et cotes de chaque référence, lues dans
                                le catalogue PDF (parser-catalogue-buronomic.py)

Deux choses viennent du catalogue et jamais du tarif : le numéro de page et
les cotes. Le tarif écrit ses dimensions dans la désignation, de façon
partielle — il ignore presque toujours la hauteur. Le catalogue les donne en
clair dans ses tableaux. On garde le tarif en second recours.

Sortie :
    catalogue-buronomic.xlsx

Usage :
    pip install openpyxl
    python generer-buronomic.py
"""

import json
import re
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from rationalisation import finition_retiree

RACINE = Path(__file__).parent
TARIF = RACINE / "buronomic.xlsx"
NOMS = RACINE / "noms-produits.json"
CATALOGUE = RACINE / "catalogue-pages-cotes.json"
SORTIE = RACINE / "catalogue-buronomic.xlsx"


COLONNES = [
    ("Gamme", 18), ("Nature", 11), ("Sur devis", 10),
    ("Catégorie", 22), ("Sous-catégorie 1", 28), ("Sous-catégorie 2", 24),
    ("Nom sur le site", 46), ("Désignation fournisseur", 50),
    ("Page cat.", 10),
    ("Réf. produit", 13), ("Réf. finition", 30), ("Réf. complète", 15),
    ("Tarif HT", 10), ("Écotaxe", 9),
    ("Description générale", 46), ("Description technique", 52),
    ("Options (réf.)", 26),
    ("Larg. min", 10), ("Larg. max", 10),
    ("Prof. min", 10), ("Prof. max", 10),
    ("Haut. min", 10), ("Haut. max", 10),
    ("Poids (kg)", 10), ("Code EAN", 15),
]

# Colonnes dont le contenu tient sur plusieurs lignes.
MULTILIGNE = (15, 16)


def cotes_du_tarif(texte):
    """Largeur, profondeur et hauteur lues dans la désignation du tarif.

    Recours de second rang, quand le catalogue est muet. Le tarif écrit les
    cotes dans le libellé, sous des formes qui varient : « PLAN DROIT L120
    x P70 », « ARMOIRE L100 x H198 », « TABLE RONDE D100 ». On lit ce qui est
    là et on laisse vide le reste : inventer une cote serait pire que de ne
    rien dire.

    Une cote peut être une plage — « PIED REGLABLE ELEC H63/128 », « PLAN
    COMPACT P80/110 ». Chaque cote sort donc en couple (mini, maxi), égaux
    quand la cote est fixe.

    Un diamètre vaut largeur et profondeur.
    """
    if not texte:
        return {}
    t = str(texte).upper().replace(",", ".")

    def cm(v):
        # Au-delà de 400, la cote est en millimètres.
        v = float(v)
        return round(v / 10, 1) if v > 400 else v

    def lire(prefixe):
        m = re.search(prefixe + r"\s*(\d{2,4}(?:\.\d)?)\s*(?:/\s*(\d{2,4}(?:\.\d)?))?", t)
        if not m:
            return None
        mini = cm(m.group(1))
        maxi = cm(m.group(2)) if m.group(2) else mini
        return (min(mini, maxi), max(mini, maxi))

    out = {}
    for cle, motif in (("L", r"\bL"), ("P", r"\bP"), ("H", r"\bH")):
        v = lire(motif)
        if v is not None:
            out[cle] = v
    if "L" not in out:
        diam = lire(r"\b[DØ]")
        if diam:
            out["L"] = out["P"] = diam
    return out


def bornes(cotes, cle):
    """Le couple (mini, maxi) d'une cote, prêt pour ses deux colonnes."""
    v = cotes.get(cle)
    return (v[0], v[1]) if v else ("", "")


def entetes(ws):
    fond = PatternFill("solid", fgColor="23262A")
    police = Font(color="FFFFFF", bold=True, size=10)
    for i, (titre, largeur) in enumerate(COLONNES, 1):
        c = ws.cell(row=1, column=i, value=titre)
        c.fill = fond
        c.font = police
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = largeur
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLONNES))}1"


def main():
    noms = json.loads(NOMS.read_text(encoding="utf-8"))
    gammes = {k: v for k, v in noms.items() if not k.startswith("_")}
    regles = noms["_rationalisation_2026"]["finitions"]
    catalogue = json.loads(CATALOGUE.read_text(encoding="utf-8"))

    # Le tarif, indexé par code racine et par gamme : un même code peut
    # servir deux gammes (la table haute ronde est chez Ensemble et chez
    # Détente), et chacune doit sortir sur ses propres lignes.
    tarif = {}
    wb = load_workbook(TARIF, read_only=True)
    rows = wb["Kits 2026"].iter_rows(values_only=True)
    next(rows)
    next(rows)
    for r in rows:
        if not r or not r[1] or not r[3]:
            continue
        tarif.setdefault((str(r[1]).strip(), str(r[3]).strip()), []).append(r)

    out = Workbook()
    ws = out.active
    ws.title = "BURONOMIC"
    entetes(ws)

    ligne = 2
    stats = {"fiches": 0, "lignes": 0, "sans_tarif": 0, "retirees": 0}
    retraits = Counter()
    cotes_source = Counter()
    manquants = []
    fiches_vides = []

    for gamme, conf in gammes.items():
        affichee = conf.get("gamme_affichee", gamme)
        cat = conf.get("categorie", "")
        sous = conf.get("sous_categories", [])
        devis = "Oui" if conf.get("sur_devis") else "Non"
        # Le texte de gamme ne sert plus que de repli : chaque fiche a le
        # sien. Cinquante-deux descriptions servies sur deux cent
        # trente-neuf fiches, c'était autant de pages en double.
        desc_gamme = conf.get("description", "")
        tech_gamme = conf.get("technique", [])
        opts = ", ".join(conf.get("options_refs", []))

        for produit in conf["produits"]:
            stats["fiches"] += 1
            nom = f"{produit['nom']} - {affichee}"
            desc = produit.get("description") or desc_gamme
            tech = "\n".join(produit.get("technique") or tech_gamme)
            premiere = True
            ecrites = 0

            # Gamme sans tarif — Essentielle, Oasys : une ligne par produit,
            # colonnes tarifaires vides.
            if not produit["refs"]:
                c = cotes_du_tarif(produit.get("detail", "") + " " + produit["nom"])
                vals = [
                    gamme, produit["nature"], devis, cat,
                    sous[0] if sous else "", sous[1] if len(sous) > 1 else "",
                    nom, produit.get("detail", ""), produit.get("page", ""),
                    "", "", "", "", "",
                    desc, tech, opts,
                    *bornes(c, "L"), *bornes(c, "P"), *bornes(c, "H"),
                    "", "",
                ]
                for j, v in enumerate(vals, 1):
                    ws.cell(row=ligne, column=j, value=v)
                ligne += 1
                stats["lignes"] += 1
                stats["sans_tarif"] += 1
                continue

            for ref in produit["refs"]:
                lignes_tarif = tarif.get((gamme, ref), [])
                if not lignes_tarif:
                    manquants.append((gamme, ref))
                    continue

                # Les cotes viennent du catalogue ; le tarif ne sert qu'à
                # combler ce que le catalogue ne dit pas. Le catalogue les
                # écrit en couple (mini, maxi) comme le tarif.
                du_cat = {k: tuple(v) for k, v in catalogue.get(ref, {}).items()
                          if k in ("L", "P", "H")}

                for r in lignes_tarif:
                    retire = finition_retiree(gamme, r[5], regles)
                    if retire:
                        retraits[retire] += 1
                        stats["retirees"] += 1
                        continue

                    # Le catalogue prime, sauf quand le tarif donne une plage
                    # là où le catalogue donne une cote fixe : « PIED
                    # REGLABLE H63/128 » en dit plus que le « H 72 cm » du
                    # bandeau de rubrique.
                    c = dict(du_cat)
                    complement = cotes_du_tarif(r[4])
                    for k, v in complement.items():
                        fixe_au_catalogue = k in c and c[k][0] == c[k][1]
                        if k not in c or (v[0] != v[1] and fixe_au_catalogue):
                            c[k] = v
                            cotes_source[f"{k} · tarif"] += 1
                        else:
                            cotes_source[f"{k} · catalogue"] += 1
                    for k in du_cat:
                        if k not in complement:
                            cotes_source[f"{k} · catalogue"] += 1

                    vals = [
                        gamme, produit["nature"], devis, cat,
                        sous[0] if sous else "", sous[1] if len(sous) > 1 else "",
                        nom, r[4], produit.get("page", ""),
                        ref, r[5], r[2], r[6], r[7],
                        desc if premiere else "",
                        tech if premiere else "",
                        opts if premiere else "",
                        *bornes(c, "L"), *bornes(c, "P"), *bornes(c, "H"),
                        r[8], r[11],
                    ]
                    for j, v in enumerate(vals, 1):
                        ws.cell(row=ligne, column=j, value=v)
                    for j in MULTILIGNE:
                        ws.cell(row=ligne, column=j).alignment = Alignment(
                            wrap_text=True, vertical="top"
                        )
                    ligne += 1
                    ecrites += 1
                    stats["lignes"] += 1
                    premiere = False

            if produit["refs"] and not ecrites:
                fiches_vides.append((gamme, produit["nom"]))

    out.save(SORTIE)

    print(f"{stats['lignes']} lignes · {stats['fiches']} fiches · {len(gammes)} gammes")
    print(f"  dont {stats['sans_tarif']} lignes sans tarif (gammes sur devis)")
    print(f"\n{stats['retirees']} lignes retirées par la rationalisation 2026 :")
    for libelle, n in retraits.most_common():
        print(f"   {n:6}  {libelle}")

    print("\nOrigine des cotes :")
    for cle in sorted(cotes_source):
        print(f"   {cotes_source[cle]:6}  {cle}")
    vides = {c: stats["lignes"] - sum(v for k, v in cotes_source.items() if k.startswith(c))
             for c in ("L", "P", "H")}
    print("   sans cote : " + " · ".join(f"{c} {n}" for c, n in vides.items()))

    if manquants:
        print(f"\n{len(manquants)} références introuvables au tarif :")
        for g, r in manquants[:10]:
            print(f"   {g} · {r}")
    if fiches_vides:
        print(f"\n{len(fiches_vides)} fiches vidées par les retraits de finition :")
        for g, n in fiches_vides:
            print(f"   {g} · {n}")

    print(f"\nÉcrit → {SORTIE.name}")


if __name__ == "__main__":
    main()
